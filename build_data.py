# -*- coding: utf-8 -*-
"""
从 Excel 源表生成 data/balcony-data.js（与页面代码完全分离的数据层）。
统一字段结构，与后续 Supabase 表结构一一对应：

  entries:  id / category / category_name / title / original_title /
            source / year / description / dimension / dimension_name /
            keywords[] / image

运行: python scripts/build_data.py
"""
import json
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_MAIN = os.path.join(ROOT, "assets", "阳台词条无图片版.xlsx")
XLSX_DIM = os.path.join(ROOT, "assets", "阳台分类_7核心维度_标签副表.xlsx")
IMG_DIR = os.path.join(ROOT, "assets", "img")
OUT = os.path.join(ROOT, "data", "balcony-data.js")


def s(v):
    return str(v).strip() if v is not None else ""


def main():
    # --- 维度定义表 ---
    wb_dim = openpyxl.load_workbook(XLSX_DIM, read_only=True)
    dimensions = []
    for r in list(wb_dim["维度说明"].iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        full = s(r[0])  # 例如 "D1 边界与过渡"
        did, _, dname = full.partition(" ")
        dimensions.append({"id": did, "name": dname, "definition": s(r[1])})

    # --- 词条 -> 维度 映射（7核心维度副表 Sheet1）---
    dim_of = {}
    for r in list(wb_dim["Sheet1"].iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        full = s(r[6])
        did, _, dname = full.partition(" ")
        dim_of[s(r[0])] = (did, dname)

    # --- 主词条表 ---
    wb = openpyxl.load_workbook(XLSX_MAIN, read_only=True)
    ws = wb["工作表1"]
    rows = list(ws.iter_rows(values_only=True))[1:]

    cat_names = {}  # 字母 -> 中文分类名
    entries = []
    for r in rows:
        if not r or not r[1]:
            continue
        cat_name, cid = s(r[0]), s(r[1])
        letter = cid.split("-")[0]
        cat_names.setdefault(letter, cat_name)

        raw_kw = s(r[7])
        # 关键词按；分隔；剔除 "A × B" 组合标签（由单词标签派生，避免检索碎片化）
        keywords = [
            k.strip()
            for k in raw_kw.replace(";", "；").split("；")
            if k.strip() and "×" not in k
        ]
        # 去重保持顺序
        keywords = list(dict.fromkeys(keywords))

        did, dname = dim_of.get(cid, ("", ""))
        img = "assets/img/%s.jpg" % cid
        if not os.path.exists(os.path.join(IMG_DIR, "%s.jpg" % cid)):
            img = None

        entries.append({
            "id": cid,
            "category": letter,
            "category_name": cat_name,
            "title": s(r[2]),
            "original_title": s(r[3]),
            "source": s(r[4]),
            "year": s(r[5]),
            "description": s(r[6]),
            "dimension": did,
            "dimension_name": dname,
            "keywords": keywords,
            "image": img,
        })

    entries.sort(key=lambda e: e["id"])

    payload = {
        "schema_version": 1,
        "generated_from": [
            "assets/阳台词条无图片版.xlsx",
            "assets/阳台分类_7核心维度_标签副表.xlsx",
        ],
        "categories": [
            {"id": k, "name": cat_names[k]} for k in sorted(cat_names)
        ],
        "dimensions": dimensions,
        "entries": entries,
    }

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("// 自动生成：scripts/build_data.py —— 请勿手改\n")
        f.write("// 数据层与页面代码分离；字段结构与 Supabase 表对应（见 js/data-source.js）\n")
        f.write("window.BALCONY_DATA = ")
        f.write(json.dumps(payload, ensure_ascii=False, indent=2))
        f.write(";\n")

    with_img = sum(1 for e in entries if e["image"])
    print("entries:", len(entries), "| with image:", with_img)
    print("categories:", cat_names)
    print("dimensions:", len(dimensions))
    print("->", OUT)


if __name__ == "__main__":
    main()
